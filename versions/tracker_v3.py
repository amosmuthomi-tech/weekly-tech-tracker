
import argparse
import warnings
import numpy as np
import pandas as pd
from scipy import stats
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

warnings.filterwarnings("ignore")


RISK_COLORS = {
    "vix":                ("C00000", "F4CCCC"),
    "10y_treasury_yield": ("375623", "E2EFDA"),
    "chicago_fed_nfci":   ("833C00", "FCE4D6"),
}
RISK_LABELS = {
    "vix":                "VIX",
    "10y_treasury_yield": "10Y Treasury Yield",
    "chicago_fed_nfci":   "Chicago Fed NFCI",
}


def thin_border():
    s = Side(style="thin", color="BFBFBF")
    return Border(left=s, right=s, top=s, bottom=s)


def autofit_columns(writer):
    for name in writer.sheets:
        ws = writer.sheets[name]
        for col in ws.columns:
            letter = get_column_letter(col[0].column)
            width = max(
                (len(str(cell.value)) for cell in col if cell.value is not None),
                default=0,
            )
            ws.column_dimensions[letter].width = min(width + 4, 45)


def load_data(csv_path, sheet_path):
    df = pd.read_excel(csv_path) if csv_path.endswith(".xlsx") else pd.read_csv(csv_path)
    grouping = pd.read_excel(sheet_path, sheet_name="grouping")
    agg_config = pd.read_excel(sheet_path, sheet_name="aggregation_config")
    print(f"[load] {len(df):,} rows | {df['symbol'].nunique()} companies | {df['series_key'].nunique()} metrics")
    return df, grouping, agg_config


def load_corr_config(sheet_path):
    cfg = pd.read_excel(sheet_path, sheet_name="correlation_config")
    print(f"[load] correlation_config: {len(cfg)} pairs")
    return cfg


def load_risk_data(risk_path):
    risk = pd.read_csv(risk_path)
    risk["date"] = pd.to_datetime(risk["date"])
    print(f"[load] risk data: {len(risk):,} rows | series: {risk['series_key'].unique().tolist()}")
    return risk


def get_group_dims(agg_config):
    dims = {}
    for col in agg_config.columns:
        if col.lower().startswith("target_"):
            dims[col[len("target_"):].title()] = col
    print(f"[config] group dimensions: {list(dims.keys())}")
    return dims


def ffill_series(df):
    df = df.copy()
    df["date"] = pd.to_datetime(df["date"])
    df = df.sort_values(["symbol", "series_key", "date"])
    df["value"] = df.groupby(["symbol", "series_key"])["value"].transform(lambda x: x.ffill())
    return df


def build_snapshot(df):
    print(f"[snapshot] nulls before ffill: {df['value'].isna().sum():,}")
    df = ffill_series(df)
    print(f"[snapshot] nulls after  ffill: {df['value'].isna().sum():,}")
    snap = (
        df.dropna(subset=["value"])
        .sort_values("date")
        .groupby(["symbol", "series_key"])
        .last()
        .reset_index()[["symbol", "series_key", "value", "date"]]
        .rename(columns={"date": "as_of_date"})
    )
    print(f"[snapshot] {len(snap):,} rows | latest: {snap['as_of_date'].max().date()}")
    return snap


def filter_target(data, group_col, target_str):
    t = target_str.strip().lower()
    if t in ("all", "nan", ""):
        return data
    allowed = [g.strip() for g in target_str.split(",")]
    return data[data[group_col].isin(allowed)]


def build_wide_snapshot(snap, grouping, agg_config, group_col, target_col):
    merged = (
        snap
        .merge(grouping[["Ticker", group_col]], left_on="symbol", right_on="Ticker", how="inner")
        .drop(columns="Ticker")
    )
    as_of = snap["as_of_date"].max().date()
    pieces = []

    for _, cfg in agg_config.iterrows():
        metric = cfg["series_key"]
        method = cfg["agg_method"].strip().lower()
        rows = filter_target(merged[merged["series_key"] == metric], group_col, str(cfg.get(target_col, "all")))
        if rows.empty:
            continue
        agg = rows.groupby(group_col)["value"].sum() if method == "sum" else rows.groupby(group_col)["value"].median()
        tmp = agg.reset_index()
        tmp.columns = [group_col, "value"]
        tmp["col_name"] = f"{metric}__{method}"
        pieces.append(tmp)

    if not pieces:
        return pd.DataFrame()

    wide = (
        pd.concat(pieces, ignore_index=True)
        .pivot(index=group_col, columns="col_name", values="value")
        .reset_index()
    )
    wide.columns.name = None
    wide.insert(0, "as_of_date", as_of)
    wide.insert(0, "group_type", group_col.lower())
    return wide


def build_full_ts(df):
    return ffill_series(df).dropna(subset=["value"])


def agg_timeseries(data, group_col, target_col, agg_config):
    pieces = []
    for _, cfg in agg_config.iterrows():
        metric = cfg["series_key"]
        method = cfg["agg_method"].strip().lower()
        rows = filter_target(data[data["series_key"] == metric], group_col, str(cfg.get(target_col, "all")))
        if rows.empty:
            continue
        agg = (
            rows.groupby(["date", group_col])["value"].sum()
            if method == "sum"
            else rows.groupby(["date", group_col])["value"].median()
        ).reset_index()
        agg["series_key"] = metric
        agg["agg_method"] = method
        pieces.append(agg)
    return pd.concat(pieces, ignore_index=True) if pieces else pd.DataFrame()


def zscore_latest(series):
    mu, sig = series.mean(), series.std()
    if sig == 0 or pd.isna(sig):
        return np.nan
    return (series.iloc[-1] - mu) / sig


def pct_latest(series):
    if len(series) < 2:
        return np.nan
    return stats.percentileofscore(series, series.iloc[-1], kind="rank")


def build_stats(ts, agg_config, group_col, latest_date):
    snap_rows, hist_rows = [], []

    for _, cfg in agg_config.iterrows():
        metric = cfg["series_key"]
        lookback = int(cfg.get("lookback_years", 3))
        do_z = bool(cfg.get("calc_z_score", False))
        do_pct = bool(cfg.get("calc_percentile", False))
        do_export = bool(cfg.get("export_history", False))

        if not any([do_z, do_pct, do_export]):
            continue

        cutoff = latest_date - pd.DateOffset(years=lookback)
        subset = ts[(ts["series_key"] == metric) & (ts["date"] >= cutoff)].copy()
        if subset.empty:
            continue

        for grp_val in subset[group_col].unique():
            grp = subset[subset[group_col] == grp_val].sort_values("date")

            if do_export:
                for _, r in grp.iterrows():
                    hist_rows.append({
                        "date": r["date"].date(),
                        group_col: grp_val,
                        "series_key": metric,
                        "agg_method": r.get("agg_method", ""),
                        "value": r["value"],
                    })

            snap_rows.append({
                group_col: grp_val,
                "series_key": metric,
                "as_of_date": grp["date"].max().date(),
                "latest_value": grp["value"].iloc[-1],
                "lookback_years": lookback,
                "lookback_mean": grp["value"].mean(),
                "lookback_std": grp["value"].std(),
                "lookback_n": len(grp),
                "z_score": zscore_latest(grp["value"]) if do_z else np.nan,
                "percentile": pct_latest(grp["value"]) if do_pct else np.nan,
            })

    return (
        pd.DataFrame(snap_rows) if snap_rows else pd.DataFrame(),
        pd.DataFrame(hist_rows) if hist_rows else pd.DataFrame(),
    )


def apply_deriv(series, deriv):
    d = deriv.strip().lower()
    if d in ("z-score", "zscore", "z_score"):
        mu, sig = series.mean(), series.std()
        if sig == 0 or pd.isna(sig):
            return pd.Series(np.nan, index=series.index)
        return (series - mu) / sig
    if d == "percentile":
        return series.rank(pct=True) * 100
    return series.copy()


def apply_calc(series, calc):
    return series.diff() if calc.strip().lower() == "delta" else series


def pearson_r(x, y):
    mask = x.notna() & y.notna()
    if mask.sum() < 5:
        return np.nan
    r, _ = stats.pearsonr(x[mask], y[mask])
    return round(float(r), 4)


def build_corr_table(ts_agg, risk_df, corr_config, group_col, target_col, latest_date):
    risk_wide = risk_df.pivot(index="date", columns="series_key", values="value").sort_index()
    records = []

    for grp_val in sorted(ts_agg[group_col].unique()):
        row = {group_col: grp_val}

        for _, cfg in corr_config.iterrows():
            corr_key = cfg["corr_key"]
            a_key = cfg["A_series_key"]
            a_deriv = str(cfg["A_series_deriv"]).strip()
            a_calc = str(cfg["A_series_calc"]).strip()
            b_key = cfg["B_series_key"]
            b_deriv = str(cfg.get("B_series_deriv", "none")).strip()
            b_calc = str(cfg["B_series_calc"]).strip()
            lookback = int(cfg["lookback_years"])

            target = str(cfg.get(target_col, "all")).strip()
            if target.lower() not in ("all", "nan", ""):
                if grp_val not in [g.strip() for g in target.split(",")]:
                    row[corr_key] = np.nan
                    continue

            cutoff = latest_date - pd.DateOffset(years=lookback)

            a_ts = (
                ts_agg[
                    (ts_agg[group_col] == grp_val)
                    & (ts_agg["series_key"] == a_key)
                    & (ts_agg["date"] >= cutoff)
                ]
                .set_index("date")["value"]
                .sort_index()
            )

            if len(a_ts) < 5 or b_key not in risk_wide.columns:
                row[corr_key] = np.nan
                continue

            b_ts = risk_wide.loc[risk_wide.index >= cutoff, b_key].dropna()
            if len(b_ts) < 5:
                row[corr_key] = np.nan
                continue

            a_d = apply_deriv(a_ts, a_deriv)
            b_d = apply_deriv(b_ts, b_deriv)

            combined = pd.concat([a_d.rename("a"), b_d.rename("b")], axis=1, join="inner").dropna()
            if len(combined) < 5:
                row[corr_key] = np.nan
                continue

            aligned = pd.DataFrame({
                "a": apply_calc(combined["a"], a_calc),
                "b": apply_calc(combined["b"], b_calc),
            }).dropna()

            row[corr_key] = pearson_r(aligned["a"], aligned["b"])

        records.append(row)

    corr_df = pd.DataFrame(records)
    col_order = [group_col] + [c for c in corr_config["corr_key"] if c in corr_df.columns]
    return corr_df[col_order]


def run_qa(snap, grouping, agg_config, v1_results):
    rows = []

    csv_syms = set(snap["symbol"].unique())
    grp_syms = set(grouping["Ticker"].unique())
    rows.append({
        "check": "company_coverage",
        "status": "PASS" if csv_syms == grp_syms else "WARN",
        "detail": f"csv: {len(csv_syms)} | grouping: {len(grp_syms)} | gap: {grp_syms - csv_syms or 'none'}",
    })

    for group_col, df in v1_results.items():
        all_g = set(grouping[group_col].unique())
        out_g = set(df[group_col].unique()) if group_col in df.columns else set()
        missing = all_g - out_g
        rows.append({
            "check": f"{group_col.lower()}_coverage",
            "status": "WARN" if missing else "PASS",
            "detail": f"all: {len(all_g)} | in output: {len(out_g)} | excluded by config: {sorted(missing) or 'none'}",
        })

    config_metrics = set(agg_config["series_key"].unique())
    csv_metrics = set(snap["series_key"].unique())
    rows.append({
        "check": "metric_coverage",
        "status": "PASS" if not (config_metrics - csv_metrics) else "WARN",
        "detail": f"config: {len(config_metrics)} | csv: {len(csv_metrics)} | missing: {config_metrics - csv_metrics or 'none'}",
    })
    rows.append({"check": "latest_date", "status": "INFO", "detail": str(snap["as_of_date"].max().date())})

    null_s = " | ".join(f"{k.lower()}_data: {v.isnull().sum().sum()} nulls" for k, v in v1_results.items())
    rows.append({"check": "nulls_in_output", "status": "INFO", "detail": null_s})

    return pd.DataFrame(rows)[["check", "status", "detail"]]


def heatmap_style(val):
    if pd.isna(val):
        return "FFFFFF", "999999", False
    v = float(val)
    if   v >=  0.70: return "C6EFCE", "276221", True
    elif v >=  0.40: return "E2EFDA", "375623", False
    elif v >=  0.10: return "F2F2F2", "595959", False
    elif v >= -0.10: return "FFFFFF", "595959", False
    elif v >= -0.40: return "FCE4D6", "843C0C", False
    elif v >= -0.70: return "F4B8A0", "843C0C", True
    else:            return "F4CCCC", "C00000", True


def make_col_label(cfg_row):
    a_tag = "z" if "z" in str(cfg_row["A_series_deriv"]).lower() else "pct"
    a_pfx = "Δ" if cfg_row["A_series_calc"] == "delta" else ""
    b_short = {"vix": "VIX", "10y_treasury_yield": "10Y", "chicago_fed_nfci": "NFCI"}
    b_lbl = b_short.get(cfg_row["B_series_key"], cfg_row["B_series_key"])
    b_pfx = "Δ" if cfg_row["B_series_calc"] == "delta" else ""
    lb = int(cfg_row["lookback_years"])
    return f"corr({a_pfx}PE {a_tag},\n{b_pfx}{b_lbl}) [{lb}Y]"


def write_corr_sheet(wb, sheet_name, corr_df, group_col, corr_config):
    if sheet_name in wb.sheetnames:
        del wb[sheet_name]
    ws = wb.create_sheet(sheet_name)

    CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
    LEFT = Alignment(horizontal="left", vertical="center")
    BD = thin_border

    data_cols = [c for c in corr_df.columns if c != group_col]
    b_map = dict(zip(corr_config["corr_key"], corr_config["B_series_key"]))
    labels = {r["corr_key"]: make_col_label(r) for _, r in corr_config.iterrows()}
    col_pos = {ck: i + 2 for i, ck in enumerate(data_cols)}

    ws.row_dimensions[1].height = 18
    ws.row_dimensions[2].height = 40
    ws.freeze_panes = "B3"

    ws["A1"].value = group_col
    ws["A1"].fill = PatternFill("solid", fgColor="1F3864")
    ws["A1"].font = Font(name="Arial", bold=True, color="FFFFFF", size=9)
    ws["A1"].alignment = CENTER
    ws["A1"].border = BD()
    ws.merge_cells("A1:A2")

    for b_key in ["vix", "10y_treasury_yield", "chicago_fed_nfci"]:
        cols_for = [ck for ck in data_cols if b_map.get(ck) == b_key]
        if not cols_for:
            continue
        c_start, c_end = col_pos[cols_for[0]], col_pos[cols_for[-1]]
        s_ltr, e_ltr = get_column_letter(c_start), get_column_letter(c_end)
        if c_start != c_end:
            ws.merge_cells(f"{s_ltr}1:{e_ltr}1")
        cell = ws[f"{s_ltr}1"]
        cell.value = RISK_LABELS[b_key]
        cell.fill = PatternFill("solid", fgColor=RISK_COLORS[b_key][0])
        cell.font = Font(name="Arial", bold=True, color="FFFFFF", size=9)
        cell.alignment = CENTER
        cell.border = BD()

    for ck in data_cols:
        b_key = b_map.get(ck, "vix")
        cell = ws[f"{get_column_letter(col_pos[ck])}2"]
        cell.value = labels.get(ck, ck)
        cell.fill = PatternFill("solid", fgColor=RISK_COLORS[b_key][0])
        cell.font = Font(name="Arial", bold=True, color="FFFFFF", size=8)
        cell.alignment = CENTER
        cell.border = BD()

    for ri, (_, drow) in enumerate(corr_df.iterrows()):
        er = ri + 3
        ws.row_dimensions[er].height = 15
        row_bg = "EBF3FF" if ri % 2 == 0 else "FFFFFF"

        gc = ws.cell(row=er, column=1)
        gc.value = drow[group_col]
        gc.font = Font(name="Arial", bold=True, size=9)
        gc.fill = PatternFill("solid", fgColor=row_bg)
        gc.alignment = LEFT
        gc.border = BD()

        for ck in data_cols:
            val = drow.get(ck)
            cell = ws.cell(row=er, column=col_pos[ck])
            bg, fg, bold = heatmap_style(val)
            cell.fill = PatternFill("solid", fgColor=bg)
            cell.border = BD()
            cell.alignment = CENTER
            if pd.notna(val):
                cell.value = float(val)
                cell.number_format = "0.00"
                cell.font = Font(name="Arial", size=9, bold=bold, color=fg)
            else:
                cell.value = "n/a"
                cell.font = Font(name="Arial", size=8, color="BFBFBF", italic=True)

    ws.column_dimensions["A"].width = 28
    for ck in data_cols:
        ws.column_dimensions[get_column_letter(col_pos[ck])].width = 13

    lr = corr_df.shape[0] + 4
    ws.row_dimensions[lr].height = 12
    ws.cell(row=lr, column=1).value = "Heatmap key:"
    ws.cell(row=lr, column=1).font = Font(name="Arial", bold=True, size=8)
    for i, (lbl, clr) in enumerate([
        ("≥0.70", "C6EFCE"), ("0.40–0.69", "E2EFDA"), ("0.10–0.39", "F2F2F2"),
        ("~0", "FFFFFF"), ("−0.10–−0.39", "FCE4D6"), ("−0.40–−0.69", "F4B8A0"), ("≤−0.70", "F4CCCC"),
    ]):
        cell = ws.cell(row=lr, column=i + 2)
        cell.value = lbl
        cell.fill = PatternFill("solid", fgColor=clr)
        cell.font = Font(name="Arial", size=8)
        cell.alignment = CENTER
        cell.border = BD()


def write_corr_config_sheet(wb, corr_config):
    name = "correlation_config"
    if name in wb.sheetnames:
        del wb[name]
    ws = wb.create_sheet(name)

    HDR = PatternFill("solid", fgColor="1F3864")
    ALT = PatternFill("solid", fgColor="EEF2F7")
    CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
    BD = thin_border

    headers = list(corr_config.columns)
    ws.append(headers)
    ws.row_dimensions[1].height = 28
    for cell in ws[1]:
        cell.fill = HDR
        cell.font = Font(name="Arial", bold=True, color="FFFFFF", size=9)
        cell.alignment = CENTER
        cell.border = BD()

    for r_idx, row in corr_config.iterrows():
        ws.append(list(row))
        for cell in ws[r_idx + 2]:
            cell.font = Font(name="Arial", size=9)
            cell.alignment = CENTER
            cell.border = BD()
            if r_idx % 2 == 1:
                cell.fill = ALT

    widths = {
        "corr_key": 22, "A_series_key": 16, "A_series_deriv": 15, "A_series_calc": 14,
        "B_series_key": 22, "B_series_deriv": 15, "B_series_calc": 14,
        "lookback_years": 14, "target_basket": 14, "target_vertical": 90,
    }
    for i, h in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(h, 15)


def run(csv_path, risk_path, sheet_path, output_path):
    df_raw, grouping, agg_config = load_data(csv_path, sheet_path)
    corr_config = load_corr_config(sheet_path)
    risk = load_risk_data(risk_path)
    group_dims = get_group_dims(agg_config)
    latest_date = pd.to_datetime(df_raw["date"].astype(str)).max()

    snap = build_snapshot(df_raw)
    ts = build_full_ts(df_raw)

    v1_out = {}
    stats_snap = {}
    stats_hist = {}
    ts_store = {}

    for group_col, target_col in group_dims.items():
        print(f"\n[v1] {group_col}: snapshot aggregation...")
        wide = build_wide_snapshot(snap, grouping, agg_config, group_col, target_col)
        if not wide.empty:
            v1_out[group_col] = wide
            print(f"     {group_col.lower()}_data: {wide.shape[0]} groups × {wide.shape[1]} cols")

        print(f"[v2] {group_col}: timeseries aggregation...")
        ts_merged = (
            ts.merge(grouping[["Ticker", group_col]], left_on="symbol", right_on="Ticker", how="inner")
            .drop(columns="Ticker")
        )
        ts_agg = agg_timeseries(ts_merged, group_col, target_col, agg_config)
        if ts_agg.empty:
            continue
        ts_store[group_col] = ts_agg

        print(f"[v2] {group_col}: z-scores and percentiles...")
        df_lat, df_hist = build_stats(ts_agg, agg_config, group_col, latest_date)
        if not df_lat.empty:
            stats_snap[group_col] = df_lat
        if not df_hist.empty:
            stats_hist[group_col] = df_hist

    print("\n[qa] running checks...")
    qa = run_qa(snap, grouping, agg_config, v1_out)
    for _, row in qa.iterrows():
        print(f"     [{row['status']}] {row['check']}: {row['detail']}")

    corr_out = {}
    for group_col, target_col in group_dims.items():
        if group_col not in ts_store:
            continue
        print(f"\n[v3] {group_col}: PE × risk correlations...")
        tbl = build_corr_table(ts_store[group_col], risk, corr_config, group_col, target_col, latest_date)
        corr_out[group_col] = tbl
        print(f"     {group_col.lower()}_correlation: {tbl.shape[0]} groups × {tbl.shape[1] - 1} corr cols")

    print(f"\n[write] {output_path}")
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        for gc, df in v1_out.items():
            df.to_excel(writer, sheet_name=f"{gc.lower()}_data", index=False)
        for gc, df in stats_snap.items():
            df.to_excel(writer, sheet_name=f"{gc.lower()}_zscore_pct", index=False)
        for gc, df in stats_hist.items():
            df.to_excel(writer, sheet_name=f"{gc.lower()}_history", index=False)
        snap.to_excel(writer, sheet_name="snapshot_debug", index=False)
        qa.to_excel(writer, sheet_name="qa_checks", index=False)
        autofit_columns(writer)

    wb = load_workbook(output_path)
    for group_col, corr_df in corr_out.items():
        write_corr_sheet(wb, f"{group_col.lower()}_correlation", corr_df, group_col, corr_config)
    write_corr_config_sheet(wb, corr_config)

    desired = (
        [f"{gc.lower()}_data" for gc in v1_out]
        + [f"{gc.lower()}_zscore_pct" for gc in stats_snap]
        + [f"{gc.lower()}_history" for gc in stats_hist]
        + [f"{gc.lower()}_correlation" for gc in corr_out]
        + ["correlation_config", "snapshot_debug", "qa_checks"]
    )
    existing = [s for s in desired if s in wb.sheetnames]
    rest = [s for s in wb.sheetnames if s not in existing]
    wb._sheets = [wb[s] for s in existing + rest]

    wb.save(output_path)
    print(f"[done] sheets: {[s.title for s in wb.worksheets]}")

    return v1_out, stats_snap, stats_hist, corr_out


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--csv",    required=True)
    parser.add_argument("--risk",   required=True)
    parser.add_argument("--sheet",  required=True)
    parser.add_argument("--output", default="tracker_v3_output.xlsx")
    args = parser.parse_args()
    run(args.csv, args.risk, args.sheet, args.output)
